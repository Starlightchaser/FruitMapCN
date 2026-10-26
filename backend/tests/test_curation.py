import json
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.curation import clean_county, curate, province_name, town_origin
from app.config import ROOT, Settings
from app.db import connect, initialize
from app.importers.excel_importer import import_workbook
from app.main import create_app


@pytest.mark.parametrize('origin,county,expected',[
    ('北京市昌平区崔村镇八家村','昌平区','北京市昌平区崔村镇'),
    ('贵州省安顺市镇宁布依族苗族自治县六马镇','镇宁布依族苗族自治县','贵州省安顺市镇宁布依族苗族自治县六马镇'),
    ('北京市怀柔区九渡河镇黄花镇村','怀柔区','北京市怀柔区九渡河镇'),
    ('河北省秦皇岛市抚宁区王汉沟村','抚宁区','河北省秦皇岛市抚宁区'),
    ('江西省上饶市德兴市张村乡张村村','','江西省上饶市德兴市张村乡'),
    ('浙江省杭州市余杭区闲林街道某社区','余杭区','浙江省杭州市余杭区闲林街道'),
    ('新疆生产建设兵团第二师25团二连','','新疆生产建设兵团第二师25团'),
])
def test_town_precision_without_guessing(origin,county,expected):
    assert town_origin(origin,county=county)[0]==expected


def test_contaminated_counties_and_ambiguous_provinces():
    assert clean_county('江安县阳春镇姜庙社区')=='江安县'
    assert clean_county('黔西南州晴隆县鸡场镇学官社区')=='晴隆县'
    assert clean_county('蓟州区')=='蓟州区'
    assert clean_county('肃州区')=='肃州区'
    assert clean_county('嵊州市')=='嵊州市'
    assert clean_county('区')==''
    names=['四川省','重庆市','广东省']
    assert province_name('广东','广东广州',names)=='广东省'
    assert province_name('四川等','四川等',names)=='四川等'
    assert province_name('四川、重庆','四川、重庆',names)=='四川、重庆'


def test_migration_preserves_raw_rows_and_curation_is_idempotent(database):
    with connect(database) as db:
        before=[tuple(row) for row in db.execute('SELECT * FROM fruit_records ORDER BY id')]
        db.executescript('DROP VIEW fruit_catalog; DROP TABLE fruit_curation; PRAGMA user_version=1;')
    initialize(database)
    with connect(database) as db:
        assert db.execute('PRAGMA user_version').fetchone()[0]==2
        report=curate(db)
        first=[dict(row) for row in db.execute('SELECT * FROM fruit_curation ORDER BY record_uid')]
        assert curate(db)==report
        assert first==[dict(row) for row in db.execute('SELECT * FROM fruit_curation ORDER BY record_uid')]
        assert before==[tuple(row) for row in db.execute('SELECT * FROM fruit_records ORDER BY id')]
        assert db.execute('SELECT full_origin FROM fruit_catalog WHERE id=1').fetchone()[0]=='测试省测试市甲县'


def test_same_town_duplicates_merge_and_reimport_recalculates(source,database):
    book=load_workbook(source)
    cells=[cell.value for cell in book.active[2]]
    cells[2]='测试省测试市甲县三村';cells[6]='三村'
    book.active.append(cells);book.save(source)
    assert import_workbook(source,database)['curation']['merged_duplicates']==1
    with connect(database,True) as db:
        assert db.execute('SELECT COUNT(*) FROM fruit_records').fetchone()[0]==4
        assert db.execute('SELECT COUNT(*) FROM fruit_catalog').fetchone()[0]==3
        assert db.execute('SELECT merged_count FROM fruit_catalog WHERE id=1').fetchone()[0]==2
    assert import_workbook(source,database)['unchanged']==4
    book=load_workbook(source);book.active['H5']='8–10月';book.active['I5']='8,9,10';book.save(source)
    assert import_workbook(source,database)['curation'].get('merged_duplicates',0)==0
    with connect(database,True) as db:
        assert db.execute('SELECT COUNT(*) FROM fruit_catalog').fetchone()[0]==4


def test_evidence_scope_filters_and_failed_curation_rollback(database,tmp_path):
    with connect(database,True) as db:
        uid=db.execute('SELECT record_uid FROM fruit_records WHERE id=1').fetchone()[0]
    evidence=tmp_path/'evidence.json'
    facts={'records':[{'record_uid':uid,'expected_product_name':'测试苹果','variety_name':'测试品种',
        'variety_scope':'产区参考','variety_source_name':'测试研究单位','variety_source_url':'https://example.org/variety',
        'taste_profile':'脆爽酸甜','taste_tags':['酸甜','脆爽'],'taste_scope':'品种参考，非镇级实测',
        'taste_source_name':'测试研究单位','taste_source_url':'https://example.org/taste','checked_at':'2026-10-26'}]}
    evidence.write_text(json.dumps(facts,ensure_ascii=False),'utf8')
    with connect(database) as db:curate(db,evidence)
    with TestClient(create_app(Settings(database_url='sqlite:///'+str(database),frontend_dist=tmp_path/'dist'))) as client:
        for value in ['documented','酸甜','脆爽']:
            assert client.get('/api/records',params={'taste':value}).json()['total']==1
            assert client.get('/api/summary',params={'taste':value}).json()['total']==1
            assert client.get('/api/map-points',params={'taste':value}).json()['total_count']==1
        detail=client.get('/api/records/1').json()
        assert detail['taste_scope']=='品种参考，非镇级实测' and detail['month_status']=='estimated'
        assert client.get('/api/records?q=测试品种').json()['total']==1
        assert client.get('/api/records?q=一村').json()['total']==0
        assert client.get('/api/records?taste=脆爽').json()['items'][0]['taste_profile']=='脆爽酸甜'
        assert client.get('/api/records?taste=无依据').status_code==422
    facts['records'][0]['taste_source_url']=''
    evidence.write_text(json.dumps(facts),'utf8')
    with pytest.raises(ValueError,match='Missing evidence'):
        with connect(database) as db:curate(db,evidence)
    with connect(database,True) as db:
        assert db.execute('SELECT taste_profile FROM fruit_catalog WHERE id=1').fetchone()[0]=='脆爽酸甜'


def test_final_source_cleanup_and_no_unjustified_bulk_taste(tmp_path):
    source=ROOT/'docs/国产水果产地月份数据库_v1.xlsx'
    if not source.exists():pytest.skip('Private source workbook not present')
    path=tmp_path/'final.db'
    report=import_workbook(source,path)
    assert report['curation']['source_records']==441
    assert report['curation']['catalog_records']==430
    with connect(path,True) as db:
        assert db.execute("SELECT COUNT(*) FROM fruit_catalog WHERE full_origin LIKE '%村' OR full_origin LIKE '%社区'").fetchone()[0]==0
        assert db.execute("SELECT COUNT(*) FROM fruit_records WHERE full_origin LIKE '%村'").fetchone()[0]>200
        assert db.execute("SELECT COUNT(*) FROM fruit_catalog WHERE product_name IN ('苹果','梨','草莓') AND taste_profile!=''").fetchone()[0]==0
        assert db.execute("SELECT COUNT(*) FROM fruit_catalog WHERE taste_profile!='' AND (taste_scope='' OR taste_source_url='')").fetchone()[0]==0
        assert db.execute("SELECT month_status FROM fruit_catalog WHERE product_name='秦脆'").fetchone()[0]=='sourced'
        assert db.execute("SELECT county FROM fruit_catalog WHERE product_name='秦脆'").fetchone()[0]==''
