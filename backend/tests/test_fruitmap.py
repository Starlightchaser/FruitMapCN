import csv
import hashlib
import json
import sqlite3
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from app.auth import hash_password
from app.config import Settings
from app.config import ROOT
from app.db import connect
from app.importers.excel_importer import import_workbook
from app.importers.geo_importer import import_locations
from app.main import create_app
from app.months import parse_months, month_status



def test_project_paths_stay_inside_fruitmap():
    expected=Path(__file__).resolve().parents[2]
    assert ROOT==expected
    assert Settings(database_url='sqlite:///backend/data/fruit.db').db_path==expected/'backend'/'data'/'fruit.db'
    assert Settings().frontend_dist==expected/'frontend'/'dist'


def test_datav_source_and_api(database,tmp_path):
    source=ROOT/'backend'/'assets'/'maps'/'datav-china.geojson'
    metadata=json.loads(source.with_suffix('.source.json').read_text(encoding='utf-8'))
    raw=source.read_bytes();boundary=json.loads(raw)
    assert hashlib.sha256(raw).hexdigest()==metadata['sha256']
    names=[feature['properties'].get('name','') for feature in boundary['features']]
    assert len(boundary['features'])==35
    assert {'台湾省','香港特别行政区','澳门特别行政区','海南省'}.issubset(names)
    assert metadata['geometry_modified'] is False
    settings=Settings(database_url='sqlite:///'+str(database),map_provider='datav',frontend_dist=tmp_path/'dist')
    assert settings.map_configured
    with TestClient(create_app(settings)) as client:
        response=client.get('/api/map-boundary')
        assert response.status_code==200
        assert response.content==raw
        config=client.get('/api/config').json()['map']
        assert config['provider']=='datav' and config['js_key']==''
    with pytest.raises(ValueError,match='Production map'):
        Settings(env='production',map_provider='datav',map_publication_approved=False).validate()


@pytest.mark.parametrize('label,expected',[
    ('9–3月（跨年）',[1,2,3,9,10,11,12]),('全年',list(range(1,13))),('未知',[]),('待补',[]),
    ('1-6月、8-12月',[1,2,3,4,5,6,8,9,10,11,12]),('6月',[6]),('12月-2月',[1,2,12])])
def test_month_parser(label,expected):
    assert parse_months(label)==expected


@pytest.mark.parametrize('label',['13月','0-3月','春夏','1至15月'])
def test_invalid_months(label):
    with pytest.raises(ValueError):parse_months(label)


def test_status_not_upgraded():
    assert month_status('沿用已有来源，未核验',[1])=='sourced'
    assert month_status('估算，待核验',[1])=='estimated'
    assert month_status('已核验',[1])=='verified'
    assert month_status('已核验',[])=='pending'


def test_dry_run_does_not_create_database(source,tmp_path):
    db=tmp_path/'not-created.db';report=import_workbook(source,db,True)
    assert report['insert']==3 and not report['committed']
    assert not db.exists()


def test_repeat_and_update(source,database):
    assert import_workbook(source,database)['unchanged']==3
    with connect(database,True) as db:
        before=db.execute('SELECT id,updated_at FROM fruit_records ORDER BY id').fetchall()
    report=import_workbook(source,database,True)
    assert report['insert']==0 and report['unchanged']==3
    book=load_workbook(source);book.active['H2']='8-10月';book.active['I2']='8,9,10';book.save(source)
    report=import_workbook(source,database)
    assert report['insert']==0 and report['update']==1
    with connect(database,True) as db:
        after=db.execute('SELECT id,updated_at FROM fruit_records ORDER BY id').fetchall()
        assert [r['id'] for r in before]==[r['id'] for r in after]
        assert before[1]['updated_at']==after[1]['updated_at']
        assert [r[0] for r in db.execute('SELECT month FROM fruit_record_months WHERE record_id=1 ORDER BY month')]==[8,9,10]


def test_invalid_import_rolls_back_everything(source,database):
    book=load_workbook(source);book.active['I2']='1,2';book.save(source)
    report=import_workbook(source,database)
    assert report['error']==1 and not report['committed']
    with connect(database,True) as db:assert db.execute('SELECT COUNT(*) FROM fruit_records').fetchone()[0]==3


def test_duplicate_uid_rejected(source,tmp_path):
    book=load_workbook(source);book.active.append([cell.value for cell in book.active[2]]);book.save(source)
    assert import_workbook(source,tmp_path/'db')['error']==1
    assert not (tmp_path/'db').exists()


@pytest.mark.parametrize('params,total',[
    ({'month':10},2),({'month':6},1),({'county':'甲县'},2),({'category':'苹果'},1),
    ({'month_status':'estimated'},1),({'q':'测试苹果'},1),({'q':"' OR 1=1 --"},0),({'q':'%'},0),
    ({'province':'另一省'},1),({'month_status':'pending'},1),({'month':10,'month_status':'pending'},0)])
def test_filters(client,params,total):
    response=client.get('/api/records',params=params)
    assert response.status_code==200
    assert response.json()['total']==total
    assert client.get('/api/summary',params=params).json()['total']==total


def test_pagination_and_detail(client):
    first=client.get('/api/records?page=1&page_size=1').json();second=client.get('/api/records?page=2&page_size=1').json()
    assert len(first['items'])==1 and first['items'][0]['id']!=second['items'][0]['id']
    assert first['total_pages']==3
    assert client.get('/api/records?page_size=9999').json()['page_size']==30
    item=client.get('/api/records/1').json()
    assert item['months']==[1,2,3,9,10,11,12]
    assert item['month_source_url']=='' and item['origin_source_url']=='https://example.org/origin'
    assert client.get('/api/records/999').status_code==404
    assert client.get('/api/records?month=13').status_code==422
    assert client.get('/api/records?page=0').status_code==422
    assert client.get('/api/records?month_status=fake').status_code==422


def test_no_static_database_or_export(client):
    for path in ['/api/export','/api/all','/api/database','/fruit.db','/backend/data/fruit.db','/data/imports/test.xlsx','/docs/test.xlsx','/.env','/source.csv','/master.json']:
        assert client.get(path).status_code==404,path
    assert 'content_hash' not in client.get('/api/records/1').json()


def test_geography_exact_match_and_aggregation(client,database,tmp_path):
    path=tmp_path/'locations.csv'
    with path.open('w',encoding='utf-8',newline='') as handle:
        writer=csv.writer(handle);writer.writerow(['county_code','province','city','county','longitude','latitude','coordinate_system','source_url','verified_at','verified'])
        writer.writerow(['999999','测试省','测试市','甲县',100,30,'CGCS2000','https://example.org/verified','2026-10-23','true'])
    assert import_locations(path,database)['resolved_records']==2
    response=client.get('/api/map-points').json()
    assert response['unresolved_count']==1
    assert response['points'][0]['record_count']==2
    assert response['points'][0]['estimated_count']==1
    assert not {'product_name','items','products','full_origin'}.intersection(response['points'][0])
    assert client.get('/api/map-points?month=6').json()['points'][0]['record_count']==1
    assert len(client.get('/api/filters').json()['counties'])==2
    assert len(client.get('/api/filters?province=另一省').json()['counties'])==1


def test_province_bubbles_include_unresolved_records_and_follow_filters(client):
    result=client.get('/api/map-points').json()
    assert result['points']==[] and result['unresolved_count']==3
    assert result['total_count']==3
    assert result['province_aggregates']==[
        dict(province='测试省',record_count=2,estimated_count=1,sourced_count=1,verified_count=0,pending_count=0),
        dict(province='另一省',record_count=1,estimated_count=0,sourced_count=0,verified_count=0,pending_count=1)]
    assert not {'product_name','items','products','full_origin'}.intersection(result['province_aggregates'][0])
    for params,total in [({'month':6},1),({'month':10},2),({'month_status':'pending'},1),
                         ({'province':'另一省'},1),({'county':'甲县','q':'测试苹果'},1),({'q':'无匹配'},0)]:
        filtered=client.get('/api/map-points',params=params).json()
        assert filtered['total_count']==total
        assert sum(row['record_count'] for row in filtered['province_aggregates'])==total
        assert client.get('/api/summary',params=params).json()['total']==total
        assert all(sum(row[key] for key in ('estimated_count','sourced_count','verified_count','pending_count'))==row['record_count']
                   for row in filtered['province_aggregates'])


def test_private_mode(database,tmp_path):
    secret='x'*40
    cfg=Settings(database_url='sqlite:///'+str(database),private_mode=True,session_secret=secret,password_hash=hash_password('test password 123456'),frontend_dist=tmp_path/'dist')
    with TestClient(create_app(cfg)) as client:
        assert client.get('/api/records').status_code==401
        assert client.get('/api/config').status_code==401
        assert not client.get('/api/auth/status').json()['authenticated']
        assert client.post('/api/auth/login',json={'password':'wrong'}).status_code==401
        login=client.post('/api/auth/login',json={'password':'test password 123456'})
        assert 'HttpOnly' in login.headers['set-cookie']
        assert client.get('/api/records').status_code==200
        assert client.post('/api/auth/logout',headers={'Origin':'https://malicious.example'}).status_code==403
        client.post('/api/auth/logout')
        assert client.get('/api/records').status_code==401


def test_limit_and_production_guard(database,tmp_path):
    with TestClient(create_app(Settings(database_url='sqlite:///'+str(database),rate_limit=2,frontend_dist=tmp_path/'dist'))) as client:
        assert client.get('/api/records').status_code==200
        assert client.get('/api/records').status_code==200
        assert client.get('/api/records').status_code==429
    with pytest.raises(ValueError,match='Production map'):
        Settings(env='production',map_provider='disabled').validate()


def test_final_workbook_integrity(tmp_path):
    workbook=Path(__file__).resolve().parents[2]/'docs'/'国产水果产地月份数据库_v1.xlsx'
    if not workbook.exists():pytest.skip('Private source workbook not present')
    database=tmp_path/'final.db';report=import_workbook(workbook,database)
    assert report['total']==441 and report['error']==0 and report['insert']==441
    assert import_workbook(workbook,database)['unchanged']==441
    with connect(database,True) as db:
        statuses=dict(db.execute('SELECT month_status,COUNT(*) FROM fruit_records GROUP BY month_status'))
        assert statuses=={'estimated':366,'sourced':64,'pending':11}
        assert db.execute("SELECT COUNT(*) FROM fruit_record_months m JOIN fruit_records r ON r.id=m.record_id WHERE r.month_status='pending'").fetchone()[0]==0
