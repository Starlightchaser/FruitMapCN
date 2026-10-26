"""Only the final workbook's 水果产地主库 is imported; no HTML/JSON fallback."""
import argparse
import hashlib
import json
import re
import sqlite3
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import load_workbook
from app.config import Settings
from app.db import initialize, connect
from app.months import parse_months, month_status
from app.curation import curate

SHEET = '水果产地主库'
REQUIRED = ('水果类目', '品种/地方产品名', '产地完整名称', '上市月份区间', '月份数字（筛选用）', '月份依据类型')


def text(value) -> str:
    return str(value).strip() if value is not None else ''


def normalized(value: str) -> str:
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', value))


def normalize_row(row: dict) -> dict:
    category, product, origin = (text(row.get(k)) for k in REQUIRED[:3])
    if not all((category, product, origin)):
        raise ValueError('水果类目、产品名、完整产地不能为空')
    label, basis = text(row.get('上市月份区间')), text(row.get('月份依据类型'))
    months = parse_months(label)
    number_months = parse_months(text(row.get('月份数字（筛选用）')))
    if months != number_months:
        raise ValueError(f'月份区间与筛选月份不一致：{label} / {number_months}')
    status = month_status(basis, months)
    if months and not basis:
        raise ValueError('月份已有值，但缺少月份依据类型')
    if '待补' in basis and months:
        raise ValueError('月份依据标为待补，但月份已有值，请人工核对')
    # Do not mislabel an origin-only official source as evidence for estimated months.
    source_name = text(row.get('产地-品种来源'))
    source_url = text(row.get('来源URL'))
    uid = text(row.get('记录UID')) or hashlib.sha256((normalized(product) + '\0' + normalized(origin)).encode()).hexdigest()[:32]
    record = dict(record_uid=uid, category=category, product_name=product, full_origin=origin,
        province=text(row.get('省份')), city=text(row.get('市/州')), county=text(row.get('县/区/旗（初步解析）')),
        township_village=text(row.get('乡镇/村（原文）')), month_range_label=label or '月份待补', month_status=status,
        month_note=text(row.get('月份说明')) or basis or ('缺少可用月份，等待补充。' if status == 'pending' else '沿用来源记录，尚待进一步核验。'),
        origin_source_name=source_name, origin_source_url=source_url,
        month_source_name=text(row.get('月份来源名称')) or (source_name if status in ('verified', 'sourced') else ''),
        month_source_url=text(row.get('月份来源URL')) or (source_url if status in ('verified', 'sourced') else ''),
        source_status=text(row.get('记录状态')), source_note=text(row.get('备注')), months=months)
    record['content_hash'] = hashlib.sha256(json.dumps(record, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    return record


def import_workbook(path: Path, db_path: Path, dry_run: bool = False) -> dict:
    report = dict(total=0, valid=0, insert=0, update=0, unchanged=0, warning=0, error=0, warnings=[], errors=[], committed=False)
    records, seen = [], set()
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET not in book.sheetnames:
            raise ValueError(f'缺少工作表：{SHEET}')
        values = iter(book[SHEET].values)
        headers = [text(x) for x in next(values)]
        if len(headers) != len(set(headers)) or not set(REQUIRED).issubset(headers):
            raise ValueError('列名重复或缺少必要列：' + ', '.join(set(REQUIRED) - set(headers)))
        for index, cells in enumerate(values, start=2):
            if not any(x is not None for x in cells):
                continue
            report['total'] += 1
            try:
                record = normalize_row(dict(zip(headers, cells)))
                if record['record_uid'] in seen:
                    raise ValueError('同一产品与产地重复，或记录UID重复；请合并来源后重新导入')
                seen.add(record['record_uid'])
                records.append(record)
                report['valid'] += 1
                if record['month_status'] == 'estimated':
                    report['warnings'].append({'row': index, 'code': 'estimated_month', 'message': '月份为估算，未升级核验状态'})
            except ValueError as exc:
                report['errors'].append({'row': index, 'message': str(exc)})
    finally:
        book.close()
    # An in-memory copy makes dry-run genuinely read-only, including no WAL side effects.
    memory = sqlite3.connect(':memory:')
    memory.row_factory = sqlite3.Row
    if db_path.exists():
        with sqlite3.connect(db_path.as_uri() + '?mode=ro', uri=True) as original:
            original.backup(memory)
    else:
        memory.executescript(Path(__file__).parents[1].joinpath('schema.sql').read_text(encoding='utf-8'))
    for record in records:
        old = memory.execute('SELECT content_hash FROM fruit_records WHERE record_uid=?', (record['record_uid'],)).fetchone()
        report['insert' if old is None else 'unchanged' if old['content_hash'] == record['content_hash'] else 'update'] += 1
    report['warning'], report['error'] = len(report['warnings']), len(report['errors'])
    memory.close()
    if dry_run or report['error']:
        return report
    initialize(db_path)
    with connect(db_path) as db:
        db.execute('BEGIN IMMEDIATE')
        now = datetime.now(timezone.utc).isoformat()
        for record in records:
            old = db.execute('SELECT id,content_hash FROM fruit_records WHERE record_uid=?', (record['record_uid'],)).fetchone()
            if old and old['content_hash'] == record['content_hash']:
                continue
            values = {k: v for k, v in record.items() if k != 'months'}
            values['updated_at'] = now
            columns = ','.join(values)
            update = ','.join(f'{k}=excluded.{k}' for k in values if k != 'record_uid')
            db.execute(f'INSERT INTO fruit_records ({columns}) VALUES ({",".join("?" for _ in values)}) ON CONFLICT(record_uid) DO UPDATE SET {update}', tuple(values.values()))
            record_id = db.execute('SELECT id FROM fruit_records WHERE record_uid=?', (record['record_uid'],)).fetchone()['id']
            db.execute('DELETE FROM fruit_record_months WHERE record_id=?', (record_id,))
            db.executemany('INSERT INTO fruit_record_months VALUES (?,?)', [(record_id, m) for m in record['months']])
        # Refresh trusted exact matches. An absent or ambiguous county stays unresolved.
        db.execute("UPDATE fruit_records SET county_code=NULL,longitude=NULL,latitude=NULL,coordinate_system=NULL,geo_status='unresolved'")
        db.execute('''UPDATE fruit_records SET
            county_code=c.county_code, longitude=c.longitude, latitude=c.latitude,
            coordinate_system=c.coordinate_system,geo_status='resolved'
            FROM county_locations c WHERE fruit_records.province=c.province AND fruit_records.county=c.county
            AND instr(fruit_records.full_origin,c.province)>0 AND instr(fruit_records.full_origin,c.county)>0''')
        report['curation'] = curate(db)
    report['committed'] = True
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('file', type=Path)
    parser.add_argument('--database', type=Path, default=Settings().db_path)
    parser.add_argument('--dry-run', action='store_true')
    parser.add_argument('--report', type=Path)
    args = parser.parse_args()
    report = import_workbook(args.file.resolve(), args.database.resolve(), args.dry_run)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k:v for k,v in report.items() if k not in ('warnings','errors')} | {'errors': report['errors']}, ensure_ascii=False, indent=2))
    raise SystemExit(1 if report['error'] else 0)


if __name__ == '__main__':
    main()
